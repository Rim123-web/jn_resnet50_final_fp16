#!/usr/bin/env python3
import csv
import ctypes
import ctypes.util
import json
import math
import os
import re
import shutil
import statistics
import subprocess
import sys
import threading
import time
import traceback
import zipfile
from pathlib import Path

ROOT = Path("/home/jetson/Downloads/classification_jetson")
MODEL_DIR = ROOT / "resnet50_inference_models"
TEST_DIR = ROOT / "test_imgs"
OUT = ROOT / "resnet50_fp32_core4_benchmark"
DIRS = ["scripts", "logs", "metrics", "resources", "reports", "exports", "raw_outputs"]
PT = MODEL_DIR / "resnet50_inference_only.pt"
ONNX = MODEL_DIR / "resnet50_inference.onnx"
TS = OUT / "exports/resnet50_fp32.torchscript.pt"
TRT_ENGINE = OUT / "exports/resnet50_fp32.engine"
IMG_EXTS = {".jpg", ".jpeg", ".png", ".bmp", ".webp", ".tif", ".tiff"}
INPUT_SHAPE = (1, 3, 640, 640)
USER_SITE = "/home/jetson/.local/lib/python3.8/site-packages"
PYTHON = sys.executable


def ensure_dirs():
    for d in DIRS:
        (OUT / d).mkdir(parents=True, exist_ok=True)


def write_json(path, obj):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(obj, indent=2), encoding="utf-8")


def run(cmd, cwd=ROOT, timeout=None):
    p = subprocess.run(cmd, shell=True, cwd=str(cwd), text=True, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, timeout=timeout)
    return p.returncode, p.stdout


def stop_stale():
    rc, out = run("ps aux | grep -E 'resnet50_fp32_core4_benchmark|trtexec' | grep -v grep || true")
    (OUT / "logs/stale_process_check_before.txt").write_text(out, encoding="utf-8")
    killed = []
    me = os.getpid()
    for line in out.splitlines():
        parts = line.split(None, 10)
        if len(parts) < 11:
            continue
        try:
            pid = int(parts[1])
        except Exception:
            continue
        cmd = parts[10]
        if pid == me or "run_resnet50_fp32_core4_benchmark.py" in cmd:
            continue
        if "resnet50_fp32_core4_benchmark" in cmd or ("trtexec" in cmd and "resnet50" in cmd):
            try:
                os.kill(pid, 15)
                killed.append({"pid": pid, "cmd": cmd})
            except Exception as exc:
                killed.append({"pid": pid, "cmd": cmd, "error": repr(exc)})
    write_json(OUT / "logs/stale_processes_stopped.json", killed)


def import_trt():
    if USER_SITE not in sys.path:
        sys.path.insert(0, USER_SITE)
    import tensorrt as trt
    sys.path[:] = [p for p in sys.path if p and p != USER_SITE and not p.startswith(USER_SITE + "/")]
    return trt


def discover_artifacts():
    patterns = ["*.pt", "*.pth", "*.onnx", "*.torchscript", "*.engine", "*.json", "*.yaml", "*.txt"]
    files = []
    for pat in patterns:
        files.extend(MODEL_DIR.rglob(pat))
    for p in [ROOT / "class_names.json", ROOT / "preprocessing.json"]:
        if p.exists():
            files.append(p)
    return sorted(set(str(p) for p in files))


def load_class_names_from_file():
    for p in [MODEL_DIR / "class_names.json", ROOT / "class_names.json"]:
        if p.exists():
            data = json.loads(p.read_text(encoding="utf-8"))
            if isinstance(data, dict):
                return [data[str(i)] if str(i) in data else data[i] for i in range(len(data))]
            return list(data)
    return None


def load_preprocessing():
    candidates = [MODEL_DIR / "preprocessing.json", ROOT / "preprocessing.json"]
    for p in candidates:
        if p.exists():
            return json.loads(p.read_text(encoding="utf-8")), str(p)
    return {"img_size": 640, "resize": [640, 640], "mean": [0.485, 0.456, 0.406], "std": [0.229, 0.224, 0.225], "input_shape": [1, 3, 640, 640]}, "assumed_imagenet"


def dataset_verify(class_names):
    class_dirs = sorted([p for p in TEST_DIR.iterdir() if p.is_dir()])
    if class_names:
        by_name = {p.name: p for p in class_dirs}
        ordered = [by_name[name] for name in class_names if name in by_name]
        ordered += [p for p in class_dirs if p.name not in set(class_names)]
        class_dirs = ordered
    else:
        class_names = [p.name for p in class_dirs]
    counts = []
    samples = []
    unreadable = []
    from PIL import Image
    for idx, d in enumerate(class_dirs):
        imgs = sorted([p for p in d.rglob("*") if p.is_file() and p.suffix.lower() in IMG_EXTS])
        ok = 0
        for p in imgs:
            try:
                with Image.open(p) as im:
                    im.verify()
                ok += 1
                samples.append({"path": str(p), "label": idx, "class": d.name})
            except Exception as exc:
                unreadable.append({"path": str(p), "error": repr(exc)})
        counts.append({"class_index": idx, "class_name": d.name, "image_count": len(imgs), "readable_count": ok})
    summary = {
        "test_dir": str(TEST_DIR),
        "class_names": [d.name for d in class_dirs],
        "class_counts": counts,
        "total_images": sum(c["image_count"] for c in counts),
        "total_readable_images": sum(c["readable_count"] for c in counts),
        "unreadable_images": unreadable,
        "supported_extensions": sorted(IMG_EXTS),
    }
    write_json(OUT / "metrics/dataset_verification.json", summary)
    with (OUT / "metrics/class_counts.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["class_index", "class_name", "image_count", "readable_count"])
        w.writeheader()
        w.writerows(counts)
    return samples, [d.name for d in class_dirs], summary


def build_model(num_classes, dropout, state_dict):
    import torch
    import torch.nn as nn
    from torchvision import models
    model = models.resnet50(weights=None)
    in_features = model.fc.in_features
    if dropout and float(dropout) > 0:
        model.fc = nn.Sequential(nn.Dropout(float(dropout)), nn.Linear(in_features, num_classes))
    else:
        model.fc = nn.Linear(in_features, num_classes)
    sd = state_dict
    cleaned = {}
    for k, v in sd.items():
        nk = k
        for prefix in ["module.", "model."]:
            if nk.startswith(prefix):
                nk = nk[len(prefix):]
        cleaned[nk] = v
    model.load_state_dict(cleaned, strict=True)
    return model


def load_checkpoint_model():
    import torch
    ckpt = torch.load(str(PT), map_location="cpu")
    if not isinstance(ckpt, dict) or "model_state_dict" not in ckpt:
        raise RuntimeError(f"Unsupported checkpoint type: {type(ckpt)}")
    class_names = ckpt.get("class_names") or load_class_names_from_file()
    if isinstance(class_names, dict):
        class_names = [class_names[str(i)] if str(i) in class_names else class_names[i] for i in range(len(class_names))]
    dropout = float(ckpt.get("dropout", 0.0))
    model = build_model(len(class_names), dropout, ckpt["model_state_dict"])
    model.eval()
    return model, class_names, ckpt


def preprocess_image(path, prep):
    import numpy as np
    from PIL import Image
    mean = np.array(prep.get("mean", [0.485, 0.456, 0.406]), dtype=np.float32).reshape(3, 1, 1)
    std = np.array(prep.get("std", [0.229, 0.224, 0.225]), dtype=np.float32).reshape(3, 1, 1)
    img = Image.open(path).convert("RGB").resize((640, 640), Image.BILINEAR)
    arr = np.asarray(img).astype(np.float32) / 255.0
    arr = arr.transpose(2, 0, 1)
    arr = (arr - mean) / std
    return np.ascontiguousarray(arr[None, ...].astype(np.float32))


def softmax_np(x):
    import numpy as np
    x = np.asarray(x).reshape(-1)
    x = x - np.max(x)
    e = np.exp(x)
    return e / np.sum(e)


def compute_metrics(pred_rows, class_names):
    n = len(class_names)
    cm = [[0 for _ in range(n)] for __ in range(n)]
    for r in pred_rows:
        cm[int(r["true_class_index"])][int(r["predicted_class_index"])] += 1
    per = []
    total = len(pred_rows)
    correct = sum(1 for r in pred_rows if r["correct"])
    for i, name in enumerate(class_names):
        tp = cm[i][i]
        fp = sum(cm[j][i] for j in range(n) if j != i)
        fn = sum(cm[i][j] for j in range(n) if j != i)
        support = sum(cm[i])
        prec = tp / (tp + fp) if tp + fp else 0.0
        rec = tp / (tp + fn) if tp + fn else 0.0
        f1 = 2 * prec * rec / (prec + rec) if prec + rec else 0.0
        per.append({"class_index": i, "class_name": name, "support": support, "precision": prec, "recall": rec, "f1": f1})
    macro_p = statistics.mean([x["precision"] for x in per]) if per else 0.0
    macro_r = statistics.mean([x["recall"] for x in per]) if per else 0.0
    macro_f1 = statistics.mean([x["f1"] for x in per]) if per else 0.0
    weighted = lambda key: sum(x[key] * x["support"] for x in per) / total if total else 0.0
    return {
        "accuracy": correct / total if total else 0.0,
        "correct": correct,
        "incorrect": total - correct,
        "num_images": total,
        "macro_precision": macro_p,
        "macro_recall": macro_r,
        "macro_f1": macro_f1,
        "weighted_precision": weighted("precision"),
        "weighted_recall": weighted("recall"),
        "weighted_f1": weighted("f1"),
        "per_class_metrics": per,
        "confusion_matrix": cm,
    }


def percentile(vals, p):
    vals = sorted(vals)
    if not vals:
        return None
    return vals[int(round((len(vals) - 1) * p / 100.0))]


def flatten(d, prefix=""):
    out = {}
    if isinstance(d, dict):
        for k, v in d.items():
            out.update(flatten(v, f"{prefix}.{k}" if prefix else str(k)))
    else:
        out[prefix] = d
    return out


def extract_power_w(stats):
    vals = []
    for k, v in flatten(stats).items():
        kl = k.lower()
        if not any(s in kl for s in ["power", "vdd", "pom_5v"]):
            continue
        if isinstance(v, (int, float)):
            val = float(v)
            if val > 1000:
                val /= 1000.0
            if 0 < val < 200:
                vals.append(val)
    return max(vals) if vals else None


def extract_max_numeric(stats, needles, lo=0, hi=100):
    vals = []
    for k, v in flatten(stats).items():
        if any(n in k.lower() for n in needles) and isinstance(v, (int, float)) and lo <= float(v) <= hi:
            vals.append(float(v))
    return max(vals) if vals else None


class Monitor:
    def __init__(self, fmt, interval=0.5):
        self.fmt = fmt
        self.interval = interval
        self.rows = []
        self.stop = threading.Event()
        self.thread = None
        self.method = "none"
    def start(self):
        try:
            import jtop
            self.method = "jtop"
            self.thread = threading.Thread(target=self._jtop, daemon=True)
            self.thread.start()
        except Exception:
            try:
                import psutil
                self.method = "psutil"
                self.thread = threading.Thread(target=self._psutil, daemon=True)
                self.thread.start()
            except Exception:
                self.method = "none"
    def _jtop(self):
        from jtop import jtop
        t0 = time.perf_counter()
        try:
            with jtop(interval=self.interval) as jetson:
                while not self.stop.is_set() and jetson.ok():
                    stats = dict(jetson.stats)
                    elapsed = time.perf_counter() - t0
                    self.rows.append({
                        "timestamp": time.time(),
                        "elapsed_seconds": elapsed,
                        "format": self.fmt,
                        "stage": "benchmark",
                        "cpu_percent": None,
                        "ram_used_mb": None,
                        "ram_total_mb": None,
                        "ram_percent": None,
                        "gpu_util_percent": extract_max_numeric(stats, ["gpu"], 0, 100),
                        "gpu_mem_used_mb": None,
                        "gpu_mem_total_mb": None,
                        "temperature_c": extract_max_numeric(stats, ["temp", "thermal"], 0, 120),
                        "power_watts": extract_power_w(stats),
                        "energy_joules_estimate": None,
                        "notes": json.dumps(stats, default=str),
                    })
                    time.sleep(self.interval)
        except Exception as exc:
            self.rows.append({"timestamp": time.time(), "elapsed_seconds": time.perf_counter() - t0, "format": self.fmt, "stage": "monitor_error", "notes": repr(exc)})
    def _psutil(self):
        import psutil
        t0 = time.perf_counter()
        while not self.stop.is_set():
            vm = psutil.virtual_memory()
            self.rows.append({
                "timestamp": time.time(), "elapsed_seconds": time.perf_counter() - t0, "format": self.fmt, "stage": "benchmark",
                "cpu_percent": psutil.cpu_percent(), "ram_used_mb": vm.used / 1048576, "ram_total_mb": vm.total / 1048576,
                "ram_percent": vm.percent, "gpu_util_percent": None, "gpu_mem_used_mb": None, "gpu_mem_total_mb": None,
                "temperature_c": None, "power_watts": None, "energy_joules_estimate": None, "notes": "psutil fallback",
            })
            time.sleep(self.interval)
    def close(self):
        self.stop.set()
        if self.thread:
            self.thread.join(timeout=3)
    def save(self, name, total_inferences, duration):
        csv_path = OUT / f"resources/{name}_resource_timeseries.csv"
        keys = ["timestamp", "elapsed_seconds", "format", "stage", "cpu_percent", "ram_used_mb", "ram_total_mb", "ram_percent", "gpu_util_percent", "gpu_mem_used_mb", "gpu_mem_total_mb", "temperature_c", "power_watts", "energy_joules_estimate", "notes"]
        with csv_path.open("w", newline="", encoding="utf-8") as f:
            w = csv.DictWriter(f, fieldnames=keys, extrasaction="ignore")
            w.writeheader()
            w.writerows(self.rows)
        powers = [float(r["power_watts"]) for r in self.rows if r.get("power_watts") not in [None, ""]]
        cpus = [float(r["cpu_percent"]) for r in self.rows if r.get("cpu_percent") not in [None, ""]]
        rams = [float(r["ram_used_mb"]) for r in self.rows if r.get("ram_used_mb") not in [None, ""]]
        gpus = [float(r["gpu_util_percent"]) for r in self.rows if r.get("gpu_util_percent") not in [None, ""]]
        temps = [float(r["temperature_c"]) for r in self.rows if r.get("temperature_c") not in [None, ""]]
        avg_power = statistics.mean(powers) if powers else None
        total_energy = avg_power * duration if avg_power is not None else None
        summary = {
            "avg_power_w": avg_power, "median_power_w": statistics.median(powers) if powers else None,
            "max_power_w": max(powers) if powers else None, "min_power_w": min(powers) if powers else None,
            "power_outlier_warning": "YES" if powers and max(powers) > 30 else "NO" if powers else "NA",
            "energy_total_j": total_energy, "energy_per_inference_j": total_energy / total_inferences if total_energy is not None and total_inferences else None,
            "avg_cpu_percent": statistics.mean(cpus) if cpus else None, "max_cpu_percent": max(cpus) if cpus else None,
            "avg_ram_mb": statistics.mean(rams) if rams else None, "max_ram_mb": max(rams) if rams else None,
            "avg_gpu_percent": statistics.mean(gpus) if gpus else None, "max_gpu_percent": max(gpus) if gpus else None,
            "avg_temperature_c": statistics.mean(temps) if temps else None, "max_temperature_c": max(temps) if temps else None,
            "telemetry_method": self.method,
        }
        write_json(OUT / f"resources/{name}_resource_summary.json", summary)
        return summary


class TensorRTRunner:
    def __init__(self, engine_path):
        import numpy as np
        self.np = np
        self.trt = import_trt()
        self.cudart, self.cudart_name = self.load_cudart()
        self.setup_cudart()
        logger = self.trt.Logger(self.trt.Logger.WARNING)
        runtime = self.trt.Runtime(logger)
        self.engine = runtime.deserialize_cuda_engine(Path(engine_path).read_bytes())
        if self.engine is None:
            raise RuntimeError("deserialize_cuda_engine returned None")
        self.context = self.engine.create_execution_context()
        self.bindings = [0] * self.engine.num_bindings
        self.allocs = []
        self.input_idx = None
        self.outputs = {}
        self.binding_info = []
        for i in range(self.engine.num_bindings):
            name = self.engine.get_binding_name(i)
            is_in = self.engine.binding_is_input(i)
            shape = tuple(self.engine.get_binding_shape(i))
            if is_in:
                self.input_idx = i
                if any(int(d) < 0 for d in shape):
                    self.context.set_binding_shape(i, INPUT_SHAPE)
            rt_shape = tuple(self.context.get_binding_shape(i))
            self.binding_info.append({"index": i, "name": name, "is_input": is_in, "engine_shape": shape, "runtime_shape": rt_shape, "dtype": str(self.engine.get_binding_dtype(i))})
        self.allocate()
    def load_cudart(self):
        candidates = []
        found = ctypes.util.find_library("cudart")
        if found: candidates.append(found)
        candidates += ["libcudart.so", "/usr/local/cuda/lib64/libcudart.so"]
        for c in candidates:
            try: return ctypes.CDLL(c), c
            except OSError: pass
        raise RuntimeError("Could not load cudart")
    def setup_cudart(self):
        self.cudart.cudaMalloc.argtypes = [ctypes.POINTER(ctypes.c_void_p), ctypes.c_size_t]
        self.cudart.cudaMalloc.restype = ctypes.c_int
        self.cudart.cudaFree.argtypes = [ctypes.c_void_p]
        self.cudart.cudaFree.restype = ctypes.c_int
        self.cudart.cudaMemcpy.argtypes = [ctypes.c_void_p, ctypes.c_void_p, ctypes.c_size_t, ctypes.c_int]
        self.cudart.cudaMemcpy.restype = ctypes.c_int
    def check(self, code, name):
        if code != 0: raise RuntimeError(f"{name} CUDA error {code}")
    def malloc(self, n):
        ptr = ctypes.c_void_p()
        self.check(self.cudart.cudaMalloc(ctypes.byref(ptr), int(n)), "cudaMalloc")
        self.allocs.append(ptr)
        return ptr
    def allocate(self):
        for i in range(self.engine.num_bindings):
            dtype = self.trt.nptype(self.engine.get_binding_dtype(i))
            shape = tuple(self.context.get_binding_shape(i))
            nbytes = int(self.np.prod(shape)) * self.np.dtype(dtype).itemsize
            ptr = self.malloc(nbytes)
            self.bindings[i] = int(ptr.value)
            if i == self.input_idx:
                self.input_dtype = dtype
                self.input_shape = shape
                self.input_ptr = ptr
            else:
                self.outputs[i] = self.np.empty(shape, dtype=dtype)
    def run(self, x):
        x = self.np.ascontiguousarray(x.astype(self.input_dtype, copy=False))
        if tuple(x.shape) != tuple(self.input_shape):
            raise RuntimeError(f"TensorRT input shape {x.shape} != {self.input_shape}")
        self.check(self.cudart.cudaMemcpy(self.input_ptr, ctypes.c_void_p(x.ctypes.data), x.nbytes, 1), "H2D")
        if not self.context.execute_v2(self.bindings):
            raise RuntimeError("execute_v2 returned False")
        outs = []
        for i, host in self.outputs.items():
            self.check(self.cudart.cudaMemcpy(ctypes.c_void_p(host.ctypes.data), ctypes.c_void_p(self.bindings[i]), host.nbytes, 2), "D2H")
            outs.append(host.copy())
        return outs[0]
    def close(self):
        for p in self.allocs:
            if p and p.value:
                self.cudart.cudaFree(p)


def evaluate_format(name, artifact, infer_fn, samples, class_names, prep, warmup=5):
    import numpy as np
    log_path = OUT / f"logs/{name}.log"
    pred_rows, times = [], []
    monitor = Monitor(name)
    try:
        for s in samples[:warmup]:
            x = preprocess_image(s["path"], prep)
            infer_fn(x)
        monitor.start()
        t_all0 = time.perf_counter()
        for s in samples:
            t0 = time.perf_counter()
            x = preprocess_image(s["path"], prep)
            t1 = time.perf_counter()
            logits = infer_fn(x)
            t2 = time.perf_counter()
            probs = softmax_np(logits)
            pred = int(np.argmax(probs))
            conf = float(probs[pred])
            t3 = time.perf_counter()
            pred_rows.append({
                "image_path": s["path"],
                "true_class": class_names[s["label"]],
                "true_class_index": s["label"],
                "predicted_class": class_names[pred] if pred < len(class_names) else "unknown",
                "predicted_class_index": pred,
                "confidence": conf,
                "correct": pred == s["label"],
                "latency_ms": (t2 - t1) * 1000,
            })
            times.append({"image_path": s["path"], "preprocess_ms": (t1-t0)*1000, "inference_ms": (t2-t1)*1000, "postprocess_ms": (t3-t2)*1000, "total_pipeline_ms": (t3-t0)*1000})
        duration = time.perf_counter() - t_all0
        monitor.close()
        res = monitor.save(name, len(samples), duration)
        metrics = compute_metrics(pred_rows, class_names)
        inf = [r["inference_ms"] for r in times]
        total = [r["total_pipeline_ms"] for r in times]
        timing = {
            "warmup_count": warmup,
            "measured_images": len(times),
            "latency_mean_ms": statistics.mean(inf),
            "latency_median_ms": statistics.median(inf),
            "latency_p95_ms": percentile(inf, 95),
            "latency_min_ms": min(inf),
            "latency_max_ms": max(inf),
            "fps": 1000.0 / statistics.mean(inf),
            "total_pipeline_fps": 1000.0 / statistics.mean(total),
            "preprocessing_time_mean_ms": statistics.mean([r["preprocess_ms"] for r in times]),
            "inference_time_mean_ms": statistics.mean(inf),
            "postprocessing_time_mean_ms": statistics.mean([r["postprocess_ms"] for r in times]),
        }
        with (OUT / f"metrics/{name}_predictions.csv").open("w", newline="", encoding="utf-8") as f:
            w = csv.DictWriter(f, fieldnames=["image_path", "true_class", "true_class_index", "predicted_class", "predicted_class_index", "confidence", "correct", "latency_ms"])
            w.writeheader(); w.writerows(pred_rows)
        with (OUT / f"metrics/{name}_latency_timeseries.csv").open("w", newline="", encoding="utf-8") as f:
            w = csv.DictWriter(f, fieldnames=["image_path", "preprocess_ms", "inference_ms", "postprocess_ms", "total_pipeline_ms"])
            w.writeheader(); w.writerows(times)
        write_json(OUT / f"metrics/{name}_timing_summary.json", timing)
        full = {"format": name, "precision": "FP32", "status": "SUCCESS", "artifact_path": str(artifact), "input_shape": "1x3x640x640", "dtype": "FP32", **metrics, **timing, "resource_summary": res, "error_message": "", "notes": ""}
        write_json(OUT / f"metrics/{name}_metrics.json", full)
        log_path.write_text("SUCCESS\n", encoding="utf-8")
        return full
    except Exception as exc:
        try: monitor.close()
        except Exception: pass
        err = repr(exc) + "\n" + traceback.format_exc()
        log_path.write_text(err, encoding="utf-8")
        out = {"format": name, "precision": "FP32", "status": "FAILED", "artifact_path": str(artifact), "input_shape": "1x3x640x640", "dtype": "FP32", "error_message": repr(exc), "notes": err}
        write_json(OUT / f"metrics/{name}_metrics.json", out)
        return out


def save_confusion_and_per_class(best_metrics, class_names):
    if not best_metrics or "confusion_matrix" not in best_metrics:
        return
    with (OUT / "metrics/confusion_matrix.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["true\\pred"] + class_names)
        for name, row in zip(class_names, best_metrics["confusion_matrix"]):
            w.writerow([name] + row)
    with (OUT / "metrics/per_class_metrics.csv").open("w", newline="", encoding="utf-8") as f:
        keys = ["class_index", "class_name", "support", "precision", "recall", "f1"]
        w = csv.DictWriter(f, fieldnames=keys)
        w.writeheader()
        w.writerows(best_metrics["per_class_metrics"])


def make_row(idx, label, m, artifact):
    res = m.get("resource_summary", {})
    return {
        "index": idx, "format": label, "precision": "FP32", "status": m.get("status", "FAILED"), "size_mb": Path(artifact).stat().st_size / 1048576 if artifact and Path(artifact).exists() else None,
        "accuracy": m.get("accuracy"), "macro_precision": m.get("macro_precision"), "macro_recall": m.get("macro_recall"), "macro_f1": m.get("macro_f1"),
        "weighted_precision": m.get("weighted_precision"), "weighted_recall": m.get("weighted_recall"), "weighted_f1": m.get("weighted_f1"),
        "latency_mean_ms": m.get("latency_mean_ms"), "latency_median_ms": m.get("latency_median_ms"), "latency_p95_ms": m.get("latency_p95_ms"), "fps": m.get("fps"),
        "avg_power_w": res.get("avg_power_w"), "median_power_w": res.get("median_power_w"), "max_power_w": res.get("max_power_w"), "min_power_w": res.get("min_power_w"),
        "energy_total_j": res.get("energy_total_j"), "energy_per_inference_j": res.get("energy_per_inference_j"),
        "avg_cpu_percent": res.get("avg_cpu_percent"), "avg_ram_mb": res.get("avg_ram_mb"), "avg_temperature_c": res.get("avg_temperature_c"),
        "telemetry_method": res.get("telemetry_method"), "power_outlier_warning": res.get("power_outlier_warning"),
        "input_shape": "1x3x640x640", "shape_verified": "YES" if m.get("status") == "SUCCESS" else "NO", "dtype_verified": "YES" if m.get("status") == "SUCCESS" else "NO",
        "artifact_path": str(artifact) if artifact else "", "error_message": m.get("error_message", ""), "notes": m.get("notes", ""),
    }


def build_trt_engine():
    if TRT_ENGINE.exists() and TRT_ENGINE.stat().st_size > 0:
        return {"status": "SUCCESS", "artifact": TRT_ENGINE, "notes": "existing output engine reused"}
    existing = ROOT / "models/resnet50_fp32_640x640.engine"
    if existing.exists() and existing.stat().st_size > 0:
        shutil.copy2(existing, TRT_ENGINE)
        (OUT / "logs/TensorRT_FP32_build.log").write_text(
            f"Reused existing FP32 TensorRT engine matching 1x3x640x640:\n{existing}\nCopied to:\n{TRT_ENGINE}\n",
            encoding="utf-8",
        )
        return {"status": "SUCCESS", "artifact": TRT_ENGINE, "notes": f"copied existing engine from {existing}"}
    cmd = f"/usr/src/tensorrt/bin/trtexec --onnx={ONNX} --saveEngine={TRT_ENGINE} --workspace=1024"
    rc, out = run(cmd, timeout=3600)
    (OUT / "logs/TensorRT_FP32_build.log").write_text("$ " + cmd + "\n" + out, encoding="utf-8")
    if rc != 0 or not TRT_ENGINE.exists() or TRT_ENGINE.stat().st_size == 0:
        return {"status": "FAILED", "error": f"trtexec rc={rc}", "log": out}
    return {"status": "SUCCESS", "artifact": TRT_ENGINE, "notes": "built from ONNX with trtexec FP32"}


def main():
    ensure_dirs()
    stop_stale()
    script_copy = OUT / "scripts/run_resnet50_fp32_core4_benchmark.py"
    if Path(__file__).resolve() != script_copy.resolve():
        shutil.copy2(Path(__file__), script_copy)
    model, ckpt_classes, ckpt = load_checkpoint_model()
    prep, prep_source = load_preprocessing()
    class_file_names = load_class_names_from_file()
    class_names = ckpt_classes or class_file_names
    samples, dataset_classes, ds_summary = dataset_verify(class_names)
    class_names = dataset_classes
    artifacts = discover_artifacts()
    sample_checks = [s["path"] for s in samples[:5]]
    write_json(OUT / "metrics/preprocessing_verification.json", {
        "resize_method": "PIL RGB resize to 640x640",
        "normalization_mean": prep.get("mean", [0.485, 0.456, 0.406]),
        "normalization_std": prep.get("std", [0.229, 0.224, 0.225]),
        "input_tensor_shape": [1, 3, 640, 640],
        "dtype": "float32",
        "preprocessing_source": prep_source,
        "sample_image_paths_checked": sample_checks,
    })
    env_lines = [
        f"Python executable: {sys.executable}",
        f"Python version: {sys.version}",
        f"Selected PyTorch model: {PT}",
        f"Selected ONNX model: {ONNX}",
        f"Test image folder: {TEST_DIR}",
        f"Class names: {class_names}",
        f"Class counts: {ds_summary['class_counts']}",
        f"Required input shape: 1x3x640x640",
        f"Target precision: FP32",
        f"Command line: {' '.join(sys.argv)}",
        "Artifacts:",
    ] + artifacts
    for mod in ["torch", "torchvision", "onnxruntime", "onnx", "jtop"]:
        try:
            m = __import__(mod)
            env_lines.append(f"{mod}: OK {getattr(m, '__version__', '')}")
        except Exception as exc:
            env_lines.append(f"{mod}: FAIL {repr(exc)}")
    try:
        import torch
        env_lines.append(f"CUDA available: {torch.cuda.is_available()}")
        env_lines.append(f"CUDA device: {torch.cuda.get_device_name(0) if torch.cuda.is_available() else 'NA'}")
    except Exception as exc:
        env_lines.append(f"CUDA check failed: {repr(exc)}")
    try:
        trt = import_trt()
        env_lines.append(f"TensorRT Python: OK {trt.__version__}")
    except Exception as exc:
        env_lines.append(f"TensorRT Python: FAIL {repr(exc)}")
    env_lines.append("trtexec: " + run("which /usr/src/tensorrt/bin/trtexec || which trtexec || true")[1].strip())
    env_lines.append("tegrastats: " + run("which tegrastats || true")[1].strip())
    (OUT / "reports/environment_and_artifact_report.txt").write_text("\n".join(env_lines) + "\n", encoding="utf-8")

    import torch
    device = torch.device("cuda:0" if torch.cuda.is_available() else "cpu")
    model = model.to(device).eval()
    def load_existing_metric(path):
        p = Path(path)
        if p.exists():
            try:
                return json.loads(p.read_text(encoding="utf-8"))
            except Exception:
                return None
        return None

    results = {}
    existing_pt = load_existing_metric(OUT / "metrics/PyTorch_FP32_metrics.json")
    if existing_pt and existing_pt.get("status") == "SUCCESS":
        results["PyTorch FP32"] = existing_pt
    else:
        with torch.no_grad():
            def infer_pt(x):
                t = torch.from_numpy(x).to(device)
                return model(t).detach().cpu().numpy()
            results["PyTorch FP32"] = evaluate_format("PyTorch_FP32", PT, infer_pt, samples, class_names, prep)
    dummy = torch.zeros(INPUT_SHAPE, dtype=torch.float32).to(device)
    existing_ts = load_existing_metric(OUT / "metrics/TorchScript_FP32_metrics.json")
    if existing_ts and existing_ts.get("status") == "SUCCESS":
        results["TorchScript FP32"] = existing_ts
    else:
      try:
        scripted = torch.jit.trace(model, dummy)
        scripted.save(str(TS))
        scripted.eval()
        with torch.no_grad():
            def infer_ts(x):
                t = torch.from_numpy(x).to(device)
                return scripted(t).detach().cpu().numpy()
            results["TorchScript FP32"] = evaluate_format("TorchScript_FP32", TS, infer_ts, samples, class_names, prep)
      except Exception as exc:
        err = {"format": "TorchScript FP32", "precision": "FP32", "status": "FAILED", "artifact_path": str(TS), "input_shape": "1x3x640x640", "dtype": "FP32", "error_message": repr(exc), "notes": traceback.format_exc()}
        write_json(OUT / "metrics/TorchScript_FP32_metrics.json", err)
        (OUT / "logs/TorchScript_FP32.log").write_text(err["notes"], encoding="utf-8")
        results["TorchScript FP32"] = err

    existing_onnx = load_existing_metric(OUT / "metrics/ONNX_FP32_metrics.json")
    if existing_onnx:
        results["ONNX FP32"] = existing_onnx
    else:
      try:
        import onnxruntime as ort
        sess = ort.InferenceSession(str(ONNX), providers=["CPUExecutionProvider"])
        input_name = sess.get_inputs()[0].name
        def infer_onnx(x):
            return sess.run(None, {input_name: x})[0]
        results["ONNX FP32"] = evaluate_format("ONNX_FP32", ONNX, infer_onnx, samples, class_names, prep)
      except Exception as exc:
        err = {"format": "ONNX FP32", "precision": "FP32", "status": "RUNTIME_UNAVAILABLE", "artifact_path": str(ONNX), "input_shape": "1x3x640x640", "dtype": "FP32", "error_message": repr(exc), "notes": traceback.format_exc()}
        write_json(OUT / "metrics/ONNX_FP32_metrics.json", err)
        (OUT / "logs/ONNX_FP32.log").write_text(err["notes"], encoding="utf-8")
        results["ONNX FP32"] = err

    trt_build = build_trt_engine()
    if trt_build["status"] == "SUCCESS":
        try:
            runner = TensorRTRunner(TRT_ENGINE)
            shape_ok = any(b["is_input"] and tuple(b["runtime_shape"]) == INPUT_SHAPE for b in runner.binding_info)
            write_json(OUT / "metrics/tensorrt_binding_info.json", runner.binding_info)
            if not shape_ok:
                raise RuntimeError(f"TensorRT shape not 1x3x640x640: {runner.binding_info}")
            def infer_trt(x):
                return runner.run(x)
            results["TensorRT FP32"] = evaluate_format("TensorRT_FP32", TRT_ENGINE, infer_trt, samples, class_names, prep)
            runner.close()
        except Exception as exc:
            err = {"format": "TensorRT FP32", "precision": "FP32", "status": "FAILED", "artifact_path": str(TRT_ENGINE), "input_shape": "1x3x640x640", "dtype": "FP32", "error_message": repr(exc), "notes": traceback.format_exc()}
            write_json(OUT / "metrics/TensorRT_FP32_metrics.json", err)
            (OUT / "logs/TensorRT_FP32.log").write_text(err["notes"], encoding="utf-8")
            results["TensorRT FP32"] = err
    else:
        err = {"format": "TensorRT FP32", "precision": "FP32", "status": "FAILED", "artifact_path": str(TRT_ENGINE), "input_shape": "1x3x640x640", "dtype": "FP32", "error_message": trt_build.get("error"), "notes": trt_build.get("log", "")}
        write_json(OUT / "metrics/TensorRT_FP32_metrics.json", err)
        results["TensorRT FP32"] = err

    shape_ver = []
    for fmt, m in results.items():
        shape_ver.append({"format": fmt, "requested_input_shape": "1x3x640x640", "observed_input_shape": "1x3x640x640" if m.get("status") == "SUCCESS" else None, "shape_verified": "YES" if m.get("status") == "SUCCESS" else "NO", "dtype_verified": "YES" if m.get("status") == "SUCCESS" else "NO", "verification_method": "shared preprocessing and runtime binding where available", "notes": m.get("error_message", "")})
    write_json(OUT / "metrics/input_shape_verification.json", shape_ver)

    order = [("PyTorch FP32", PT), ("TorchScript FP32", TS), ("ONNX FP32", ONNX), ("TensorRT FP32", TRT_ENGINE)]
    rows = [make_row(i+1, label, results[label], art) for i, (label, art) in enumerate(order)]
    with (OUT / "metrics/resnet50_fp32_core4_benchmark_table.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        w.writeheader(); w.writerows(rows)
    write_json(OUT / "metrics/resnet50_fp32_core4_benchmark_table.json", rows)
    best_acc = max([r for r in rows if r["status"] == "SUCCESS"], key=lambda r: r["accuracy"], default=None)
    best_f1 = max([r for r in rows if r["status"] == "SUCCESS"], key=lambda r: r["macro_f1"], default=None)
    best_fps = max([r for r in rows if r["status"] == "SUCCESS"], key=lambda r: r["fps"], default=None)
    best_energy = min([r for r in rows if r["status"] == "SUCCESS" and r.get("energy_per_inference_j") is not None], key=lambda r: r["energy_per_inference_j"], default=None)
    best_metrics = results.get(best_acc["format"]) if best_acc else None
    save_confusion_and_per_class(best_metrics, class_names)
    report = [
        "# ResNet50 FP32 Core-4 Classification Benchmark",
        "",
        "Objective: custom fixed-shape ResNet50 classification FP32 benchmark on Jetson Nano.",
        "Input shape: 1x3x640x640. FP32 only. Formats attempted: PyTorch, TorchScript, ONNX, TensorRT.",
        "No training, fine-tuning, FP16, INT8, quantization, YOLO metrics, mAP, NMS, or box decoding were used.",
        f"Selected PyTorch model: {PT}",
        f"Selected ONNX model: {ONNX}",
        f"Test image path: {TEST_DIR}",
        f"Class names: {class_names}",
        f"Total images: {len(samples)}",
        f"Preprocessing: RGB resize 640x640, divide by 255, normalize mean={prep.get('mean')}, std={prep.get('std')}, HWC to CHW, FP32 batch=1.",
        "",
        "Best FP32 format by accuracy: " + (best_acc["format"] if best_acc else "none"),
        "Best FP32 format by macro F1: " + (best_f1["format"] if best_f1 else "none"),
        "Best FP32 format by FPS: " + (best_fps["format"] if best_fps else "none"),
        "Best FP32 format by energy: " + (best_energy["format"] if best_energy else "none"),
        "",
        "Benchmark table: metrics/resnet50_fp32_core4_benchmark_table.csv",
        "Per-class metrics: metrics/per_class_metrics.csv",
        "Confusion matrix: metrics/confusion_matrix.csv",
        "",
        "Next step: compare against FP16 and INT8 deployment variants later.",
    ]
    (OUT / "reports/resnet50_fp32_core4_benchmark_report.txt").write_text("\n".join(report) + "\n", encoding="utf-8")
    (OUT / "reports/resnet50_fp32_core4_benchmark_report.md").write_text("\n".join(report) + "\n", encoding="utf-8")
    try:
        from docx import Document
        doc = Document()
        doc.add_heading("ResNet50 FP32 Core-4 Classification Benchmark", 0)
        for line in report[2:]:
            doc.add_paragraph(line)
        doc.save(OUT / "reports/resnet50_fp32_core4_benchmark_report.docx")
    except Exception as exc:
        (OUT / "logs/docx_report_skipped.log").write_text(repr(exc), encoding="utf-8")
    final = {
        "selected_pytorch_model_path": str(PT),
        "selected_onnx_model_path": str(ONNX),
        "test_image_folder": str(TEST_DIR),
        "class_names": class_names,
        "total_images_evaluated": len(samples),
        "requested_input_shape": "1x3x640x640",
        "successful_formats": [r["format"] for r in rows if r["status"] == "SUCCESS"],
        "failed_formats": [r["format"] + ":" + r["status"] for r in rows if r["status"] != "SUCCESS"],
        "best_fp32_format_by_accuracy": best_acc["format"] if best_acc else None,
        "best_fp32_format_by_macro_f1": best_f1["format"] if best_f1 else None,
        "best_fp32_format_by_fps": best_fps["format"] if best_fps else None,
        "best_fp32_format_by_energy_per_inference": best_energy["format"] if best_energy else None,
        "output_folder": str(OUT),
        "table_path": str(OUT / "metrics/resnet50_fp32_core4_benchmark_table.csv"),
        "report_path": str(OUT / "reports/resnet50_fp32_core4_benchmark_report.txt"),
        "zip_path": str(ROOT / "resnet50_fp32_core4_benchmark.zip"),
    }
    write_json(OUT / "metrics/final_summary.json", final)
    zip_path = ROOT / "resnet50_fp32_core4_benchmark.zip"
    if zip_path.exists(): zip_path.unlink()
    with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as z:
        for p in OUT.rglob("*"):
            if p.is_file():
                z.write(p, p.relative_to(OUT.parent))
    print(json.dumps(final, indent=2))
    print("RESNET50 CLASSIFICATION FIXED 640x640 FP32 CORE-4 BENCHMARK COMPLETE")


if __name__ == "__main__":
    main()
